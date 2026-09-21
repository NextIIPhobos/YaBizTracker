import tempfile, unittest
from pathlib import Path
from yabiztracker.database.database import Database
from yabiztracker.api.yandex import YandexAPI
from yabiztracker.api.errors import ApiServerError
from yabiztracker.services.export_service import ExportService
from openpyxl import load_workbook

class FakeSearch:
    def __init__(self): self.calls=[]
    def search_page(self,text,bbox,skip=0):
        self.calls.append((text,skip))
        if skip==0:
            def make_result(oid, name):
                return {"geometry": {"coordinates": [50.1, 53.2]}, "properties": {"name": name, "uri": f"https://yandex.ru/maps/org/x/{oid}", "CompanyMetaData": {"id": oid, "name": name, "Address": {"formatted": "A"}}}}
            return [make_result('1', 'One'), make_result('2', 'Two')], True
        return [{"geometry":{"coordinates":[50.2,53.3]},"properties":{"name":"Three","uri":"https://yandex.ru/maps/org/x/3","CompanyMetaData":{"id":"3","name":"Three"}}}],False


class FlakySearch(FakeSearch):
    def __init__(self): super().__init__(); self.fail_once=True
    def search_page(self,text,bbox,skip=0):
        if skip==50 and self.fail_once:
            self.calls.append((text,skip)); self.fail_once=False
            raise ApiServerError('temporary',api='search',status=503)
        return FakeSearch.search_page(self,text,bbox,skip)

class IntegrationTests(unittest.TestCase):
    def test_checkpoint_resumes_after_page_failure(self):
        with tempfile.TemporaryDirectory() as d:
            db=Database(str(Path(d)/'organizations.db')); api=YandexAPI('',''); api.search=FlakySearch(); cities=[{"name":"Samara","bbox":"50,53~50.3,53.4"}]; cats=['Квесты']
            def on_page(city,cat,skip,orgs,raw): db.upsert_organizations(orgs,city['name'])
            api.search_organizations(cities,cats,on_page=on_page)
            self.assertEqual(api.search.calls,[('Квесты',0),('Квесты',50)])
            # The failed page was not checkpointed, so the next invocation resumes at skip=50.
            api.search_organizations(cities,cats,on_page=on_page,get_start_skip=lambda c,q: db.get_checkpoint(1,c['name'],q)[0])
            self.assertEqual(api.search.calls[-1],('Квесты',50)); self.assertEqual(len(db.get_all_organizations()),3); db.close()

    def test_full_business_cycle(self):
        with tempfile.TemporaryDirectory() as d:
            db=Database(str(Path(d)/'organizations.db')); api=YandexAPI('',''); api.search=FakeSearch(); cities=[{"name":"Samara","bbox":"50,53~50.3,53.4"}]; cats=['Квесты']
            seen=[]
            def on_page(city,cat,skip,orgs,raw):
                seen.extend(orgs); db.upsert_organizations(orgs,city['name'])
            api.search_organizations(cities,cats,on_page=on_page)
            self.assertEqual(len(db.get_all_organizations()),3)
            db.update_crm('1',status='Клиент',responsible='Manager',comment='Call')
            self.assertEqual(db.get_by_id('1')['status'],'Клиент'); self.assertTrue(db.history('1'))
            out=Path(d)/'leads.xlsx'; ExportService().export(str(out),db.get_all_organizations()); self.assertTrue(out.exists()); wb=load_workbook(out,read_only=True); self.assertEqual(wb['Организации'].max_row,4); wb.close(); db.close()

if __name__=='__main__': unittest.main()

class ScanSignatureTests(unittest.TestCase):
    def test_signature_is_order_independent(self):
        from yabiztracker.domain.scan import scan_signature
        a=scan_signature([{"name":"Samara","bbox":"b"},{"name":"Tolyatti","bbox":"a"}], ["Банк","Квесты"], ["Спорт","Банк"])
        b=scan_signature([{"name":"Tolyatti","bbox":"a"},{"name":"Samara","bbox":"b"}], ["Квесты","Банк","банк"], ["банк","Спорт"])
        self.assertEqual(a,b)

    def test_search_stats_count_requests_and_exclusions_separately(self):
        api=YandexAPI('','')
        class Search:
            def search_page(self,text,bbox,skip=0):
                return [
                    {"geometry":{"coordinates":[50.1,53.2]},"properties":{"name":"Allowed","CompanyMetaData":{"id":"1","Categories":[{"name":"Квесты"}]}}},
                    {"geometry":{"coordinates":[50.1,53.2]},"properties":{"name":"Excluded","CompanyMetaData":{"id":"2","Categories":[{"name":"Квесты"},{"name":"Банк"}]}}},
                ],False
        api.search=Search()
        _,stats=api.search_organizations([{"name":"Samara","bbox":"b"}],['Квесты'],excluded_categories=['Банк'])
        self.assertEqual(stats['search_requests'],1)
        self.assertEqual(stats['api_results'],2)
        self.assertEqual(stats['excluded_by_category'],1)
        self.assertEqual(stats['duplicates'],0)
        self.assertEqual(stats['accepted_results'],1)
