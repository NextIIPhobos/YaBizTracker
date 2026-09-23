import tempfile, unittest
from pathlib import Path
from openpyxl import load_workbook
from yabiztracker.services.export_service import ExportService

class ExportTests(unittest.TestCase):
    def test_fallback_without_openpyxl(self):
        import builtins
        original=builtins.__import__
        def fake_import(name,*args,**kwargs):
            if name.startswith('openpyxl'): raise ModuleNotFoundError('simulated missing openpyxl')
            return original(name,*args,**kwargs)
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'fallback.xlsx'
            o={"name":"A","website":"https://example.com"}
            old=builtins.__import__; builtins.__import__=fake_import
            try: engine=ExportService().export(str(p),[o])
            finally: builtins.__import__=old
            self.assertEqual(engine,'fallback'); self.assertTrue(p.exists()); self.assertTrue(p.read_bytes().startswith(b'PK'))

    def test_roundtrip_and_sanitisation(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'out.xlsx'; o={"name":"A\x00B","address":"X\x0bY","city_name":"Самара","category":"C","subcategory":"S","age_days":1,"phone":"1","email":"a@b","website":"https://example.com","social_links":"{\"vk\":\"https://vk.com/a\"}","status":"Новый","comment":"C","responsible":"R","next_contact_date":"2026-09-12 12:00","score":10}
            engine=ExportService().export(str(p),[o]); self.assertTrue(p.exists()); wb=load_workbook(p); ws=wb['Организации']; self.assertEqual(ws['A2'].value,'AB'); self.assertEqual(ws['C2'].value,'Самара'); self.assertEqual(ws['H2'].hyperlink.target,'https://example.com'); wb.close(); self.assertIn(engine,('openpyxl','fallback'))

    def test_export_selected_columns_and_website_hyperlink_position(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'selected.xlsx'
            o={"name":"Компания","address":"Адрес","city_name":"Самара","category":"C","subcategory":"S",
               "age_days":2,"phone":"+7999","email":"a@b","website":"https://example.com","status":"Новый"}
            # Название, Сайт, E-mail — website moves to column B in the resulting workbook.
            engine=ExportService().export(str(p),[o],columns=[0,7,6])
            wb=load_workbook(p); ws=wb['Организации']
            self.assertEqual([ws.cell(1,c).value for c in range(1,4)],['Название','Сайт','E-mail'])
            self.assertEqual(ws['A2'].value,'Компания')
            self.assertEqual(ws['B2'].hyperlink.target,'https://example.com')
            self.assertEqual(ws['C2'].value,'a@b')
            self.assertEqual(ws.max_column,3)
            wb.close(); self.assertIn(engine,('openpyxl','fallback'))

    def test_export_rejects_empty_or_invalid_columns(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'bad.xlsx'; o={"name":"A"}
            with self.assertRaises(ValueError): ExportService().export(str(p),[o],columns=[])
            with self.assertRaises(ValueError): ExportService().export(str(p),[o],columns=[0,0])
            with self.assertRaises(ValueError): ExportService().export(str(p),[o],columns=[len(ExportService.HEADERS)])

if __name__=='__main__': unittest.main()
